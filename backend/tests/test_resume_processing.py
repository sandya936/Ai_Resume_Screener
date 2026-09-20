import io
import pytest
import docx
from httpx import AsyncClient
from app.core.config import settings


def generate_valid_pdf_bytes() -> bytes:
    return (
        b"%PDF-1.4\n"
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n"
        b"4 0 obj\n<< /Length 64 >>\nstream\n"
        b"BT /F1 12 Tf 100 700 Td (John Doe Senior Software Architect Resume) Tj ET\n"
        b"endstream\nendobj\n"
        b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
        b"xref\n0 6\n"
        b"0000000000 65535 f \n0000000009 00000 n \n0000000058 00000 n \n0000000115 00000 n \n"
        b"0000000244 00000 n \n0000000357 00000 n \n"
        b"trailer\n<< /Size 6 /Root 1 0 R >>\n"
        b"startxref\n435\n%%EOF\n"
    )


def generate_valid_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("Jane Doe - Senior Engineer", level=1)
    doc.add_paragraph("Experienced software engineer with expertise in FastAPI and Next.js.")
    doc.add_paragraph("• Architected production systems", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_upload_valid_pdf(async_client: AsyncClient):
    reg = await async_client.post(
        "/api/v1/auth/register",
        json={"email": "pdf_user@example.com", "password": "Password123!", "full_name": "PDF User"},
    )
    token = reg.json()["data"]["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    pdf_bytes = generate_valid_pdf_bytes()
    files = {"file": ("john_doe_resume.pdf", pdf_bytes, "application/pdf")}

    response = await async_client.post("/api/v1/resumes/upload", headers=headers, files=files)
    assert response.status_code == 201
    data = response.json()
    assert data["success"] is True
    assert data["data"]["processing_status"] == "completed"
    assert "John Doe" in data["data"]["extracted_text_preview"]


@pytest.mark.asyncio
async def test_upload_valid_docx(async_client: AsyncClient):
    reg = await async_client.post(
        "/api/v1/auth/register",
        json={"email": "docx_user@example.com", "password": "Password123!", "full_name": "DOCX User"},
    )
    token = reg.json()["data"]["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    docx_bytes = generate_valid_docx_bytes()
    files = {
        "file": (
            "jane_doe_resume.docx",
            docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }

    response = await async_client.post("/api/v1/resumes/upload", headers=headers, files=files)
    assert response.status_code == 201
    data = response.json()
    assert data["success"] is True
    assert "Jane Doe" in data["data"]["extracted_text_preview"]
    assert len(data["data"]["normalized_document"]["bullet_points"]) >= 1


@pytest.mark.asyncio
async def test_upload_corrupted_pdf(async_client: AsyncClient):
    reg = await async_client.post(
        "/api/v1/auth/register",
        json={"email": "corrupt_user@example.com", "password": "Password123!", "full_name": "Corrupt User"},
    )
    token = reg.json()["data"]["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    corrupt_bytes = b"This is not a PDF file header at all."
    files = {"file": ("fake.pdf", corrupt_bytes, "application/pdf")}

    response = await async_client.post("/api/v1/resumes/upload", headers=headers, files=files)
    assert response.status_code == 400
    data = response.json()
    assert data["success"] is False
    assert data["error"]["code"] == "CORRUPTED_FILE"


@pytest.mark.asyncio
async def test_upload_empty_document(async_client: AsyncClient):
    reg = await async_client.post(
        "/api/v1/auth/register",
        json={"email": "empty_user@example.com", "password": "Password123!", "full_name": "Empty User"},
    )
    token = reg.json()["data"]["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    files = {"file": ("empty.pdf", b"", "application/pdf")}

    response = await async_client.post("/api/v1/resumes/upload", headers=headers, files=files)
    assert response.status_code == 400
    data = response.json()
    assert data["success"] is False
    assert data["error"]["code"] == "EMPTY_FILE"


@pytest.mark.asyncio
async def test_upload_unsupported_format(async_client: AsyncClient):
    reg = await async_client.post(
        "/api/v1/auth/register",
        json={"email": "unsupported_user@example.com", "password": "Password123!", "full_name": "Bad Ext User"},
    )
    token = reg.json()["data"]["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    files = {"file": ("script.sh", b"#!/bin/bash\necho hello", "text/x-shellscript")}

    response = await async_client.post("/api/v1/resumes/upload", headers=headers, files=files)
    assert response.status_code == 400
    data = response.json()
    assert data["success"] is False
    assert data["error"]["code"] in ["UNSUPPORTED_EXTENSION", "UNSUPPORTED_MIME_TYPE"]


@pytest.mark.asyncio
async def test_upload_oversized_file(async_client: AsyncClient):
    reg = await async_client.post(
        "/api/v1/auth/register",
        json={"email": "large_user@example.com", "password": "Password123!", "full_name": "Large User"},
    )
    token = reg.json()["data"]["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # Generate oversized file (> 10MB)
    large_bytes = b"%PDF-1.4\n" + (b"0" * (settings.MAX_FILE_SIZE_BYTES + 1024))
    files = {"file": ("huge.pdf", large_bytes, "application/pdf")}

    response = await async_client.post("/api/v1/resumes/upload", headers=headers, files=files)
    assert response.status_code == 400
    data = response.json()
    assert data["success"] is False
    assert data["error"]["code"] == "FILE_TOO_LARGE"


@pytest.mark.asyncio
async def test_malicious_filename_sanitization(async_client: AsyncClient):
    reg = await async_client.post(
        "/api/v1/auth/register",
        json={"email": "sanitize_user@example.com", "password": "Password123!", "full_name": "Sanitize User"},
    )
    token = reg.json()["data"]["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    pdf_bytes = generate_valid_pdf_bytes()
    files = {"file": ("../../../etc/passwd.pdf", pdf_bytes, "application/pdf")}

    response = await async_client.post("/api/v1/resumes/upload", headers=headers, files=files)
    assert response.status_code == 201
    data = response.json()
    filename = data["data"]["resume"]["versions"][0]["filename"]
    assert "../" not in filename
    assert filename == "passwd.pdf"


@pytest.mark.asyncio
async def test_unauthorized_resume_access(async_client: AsyncClient):
    # User A uploads resume
    reg_a = await async_client.post(
        "/api/v1/auth/register",
        json={"email": "owner_a@example.com", "password": "Password123!", "full_name": "Owner A"},
    )
    token_a = reg_a.json()["data"]["tokens"]["access_token"]
    headers_a = {"Authorization": f"Bearer {token_a}"}

    pdf_bytes = generate_valid_pdf_bytes()
    files = {"file": ("resume_a.pdf", pdf_bytes, "application/pdf")}
    res_a = await async_client.post("/api/v1/resumes/upload", headers=headers_a, files=files)
    resume_id = res_a.json()["data"]["resume"]["id"]

    # User B tries to fetch User A's resume
    reg_b = await async_client.post(
        "/api/v1/auth/register",
        json={"email": "attacker_b@example.com", "password": "Password123!", "full_name": "Attacker B"},
    )
    token_b = reg_b.json()["data"]["tokens"]["access_token"]
    headers_b = {"Authorization": f"Bearer {token_b}"}

    get_res = await async_client.get(f"/api/v1/resumes/{resume_id}", headers=headers_b)
    assert get_res.status_code == 403
    assert get_res.json()["error"]["code"] == "FORBIDDEN"


@pytest.mark.asyncio
async def test_get_extracted_text_and_delete_resume(async_client: AsyncClient):
    reg = await async_client.post(
        "/api/v1/auth/register",
        json={"email": "full_flow@example.com", "password": "Password123!", "full_name": "Full Flow User"},
    )
    token = reg.json()["data"]["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    docx_bytes = generate_valid_docx_bytes()
    files = {
        "file": (
            "full_flow.docx",
            docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }

    upload_res = await async_client.post("/api/v1/resumes/upload", headers=headers, files=files)
    resume_id = upload_res.json()["data"]["resume"]["id"]
    version_id = upload_res.json()["data"]["version_id"]

    # 1. Fetch extracted text
    ext_res = await async_client.get(
        f"/api/v1/resumes/versions/{version_id}/extracted-text", headers=headers
    )
    assert ext_res.status_code == 200
    ext_data = ext_res.json()["data"]
    assert ext_data["status"] == "completed"
    assert "Jane Doe" in ext_data["raw_text"]

    # 2. Delete resume
    del_res = await async_client.delete(f"/api/v1/resumes/{resume_id}", headers=headers)
    assert del_res.status_code == 200
    assert del_res.json()["success"] is True

    # 3. Confirm deletion
    confirm_get = await async_client.get(f"/api/v1/resumes/{resume_id}", headers=headers)
    assert confirm_get.status_code == 404
